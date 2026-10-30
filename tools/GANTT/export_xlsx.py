#!/usr/bin/env python3
"""Build the Cool Spine GANTT workbook from roadmap.yaml.

One column per sprint, with year and quarter bands and a date header.
A task paints a green bar from its start sprint through its end sprint.
A milestone paints one blue cell on the sprint that contains its date.
Status is the roadmap status. The file is named with the day it is written:
Cool_Spine_Roadmap_YYYY-MM-DD.xlsx.

    python tools/GANTT/export_xlsx.py
"""
from __future__ import annotations

import argparse
import os
import tempfile
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from roadmap_io import RoadmapError, load_roadmap, sprint_starts

HERE = Path(__file__).resolve().parent
FONT = "Aptos Narrow"
INK = "FF1C2430"
MUTED = "FF5C6B7A"
BAR = "FF4EA72E"
MILESTONE = "FF00B0F0"
EPIC = "FFE8EEF5"
CURRENT = "FFF3F7FC"
CURRENT_HEADER = "FFE7F0FA"
HEADER = "FFF6F7F9"
LINE = "FFD5DBE3"

NAME_COL = 1
STATUS_COL = 2
FIRST_SPRINT_COL = 3
YEAR_ROW = 1
QUARTER_ROW = 2
DATE_ROW = 3
HEADER_ROW = 4

STATUS_LABELS = {
    "done": "Done",
    "in_progress": "In progress",
    "future": "Future",
}
STATUS_LIST = ("Done", "In progress", "Future", "Epic", "Milestone")
STATUS_COLORS = {
    "Done": ("FFE7F2EA", "FF1E5A38"),
    "In progress": ("FFE6EFFA", "FF1A4478"),
    "Future": ("FFEEF0F3", "FF4B5563"),
    "Milestone": ("FFFFF4E8", "FF9A3412"),
}

THIN = Side(style="thin", color=LINE)
DATE_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RULE = Border(bottom=Side(style="medium", color=INK))
CENTER = Alignment(horizontal="center", vertical="center")
BAND = Alignment(horizontal="left", vertical="center", indent=1)
NAME = Alignment(horizontal="left", vertical="center")


def build_workbook(data: dict, today: date | None = None) -> Workbook:
    """Return a workbook for a roadmap dict from load_roadmap."""
    today = today or date.today()
    step = data["timeline"]["sprint_days"]
    sprints = sprint_starts(data)
    current = _current_sprint(sprints, step, today)
    rows = _rows(data)

    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(data["title"])
    wb.properties.title = data["title"]
    wb.properties.description = "Generated from roadmap.yaml"

    _write_corner(ws, data["title"])
    _write_bands(ws, sprints)
    _write_dates(ws, sprints, current)
    _write_column_titles(ws, len(sprints))
    _write_rows(ws, rows, sprints, step, current)
    _finish_sheet(ws, len(sprints), len(rows))
    return wb


def save_workbook(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=path.stem + "_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as exc:
        raise SystemExit(f"Could not write {path}. Close it if it is open in Excel.") from exc
    finally:
        if tmp.exists():
            tmp.unlink()


def _rows(data: dict) -> list[dict]:
    rows = []
    milestones = []
    for epic in data["epics"]:
        epic_milestones = [item for item in epic["items"] if item["kind"] == "milestone"]
        milestones.extend(epic_milestones)
        if epic["id"] == "milestones":
            continue
        rows.append({"kind": "epic", "name": epic["name"], "status": "Epic"})
        for item in epic["items"]:
            if item["kind"] != "task":
                continue
            rows.append({
                "kind": "task",
                "name": item["name"],
                "status": STATUS_LABELS[item["status"]],
                "start": item.get("start"),
                "end": item.get("end"),
            })
    if milestones:
        rows.append({"kind": "epic", "name": "Milestones", "status": "Epic"})
        for item in milestones:
            rows.append({
                "kind": "milestone",
                "name": item["name"],
                "status": "Milestone",
                "date": item.get("date"),
            })
    return rows


def _write_corner(ws, title: str) -> None:
    title_cell = ws.cell(YEAR_ROW, NAME_COL, title)
    title_cell.font = _font(14, bold=True)
    title_cell.alignment = Alignment(vertical="center")
    note = ws.cell(QUARTER_ROW, NAME_COL, "Generated from roadmap.yaml")
    note.font = _font(9, color=MUTED)
    note.alignment = Alignment(vertical="center")
    ws.row_dimensions[YEAR_ROW].height = 22
    ws.column_dimensions["A"].width = 66
    ws.column_dimensions["B"].width = 16


def _write_bands(ws, sprints: list[date]) -> None:
    origin = sprints[0].year
    years = [f"Year {sprint.year - origin + 1}" for sprint in sprints]
    quarters = [f"Q{(sprint.month - 1) // 3 + 1} {sprint.year}" for sprint in sprints]
    _write_band(ws, YEAR_ROW, years, _font(bold=True))
    _write_band(ws, QUARTER_ROW, quarters, _font())


def _write_band(ws, row: int, labels: list[str], font: Font) -> None:
    for label, start, end in _runs(labels):
        column = FIRST_SPRINT_COL + start
        cell = ws.cell(row, column, label)
        cell.font = font
        cell.alignment = BAND
        if end > start:
            ws.merge_cells(
                start_row=row,
                start_column=column,
                end_row=row,
                end_column=FIRST_SPRINT_COL + end,
            )


def _write_dates(ws, sprints: list[date], current: date | None) -> None:
    ws.row_dimensions[DATE_ROW].height = 18
    for index, sprint in enumerate(sprints):
        column = FIRST_SPRINT_COL + index
        cell = ws.cell(DATE_ROW, column, sprint)
        cell.number_format = "dd-mmm-yy"
        cell.alignment = CENTER
        cell.border = DATE_BORDER
        cell.font = _font(bold=True, color="FF1A4478") if sprint == current else _font()
        if sprint == current:
            cell.fill = _fill(CURRENT_HEADER)
        ws.column_dimensions[get_column_letter(column)].width = 13


def _write_column_titles(ws, sprint_count: int) -> None:
    ws.cell(HEADER_ROW, NAME_COL, "Workstream / Task")
    status = ws.cell(HEADER_ROW, STATUS_COL, "Status")
    last = FIRST_SPRINT_COL + sprint_count - 1
    for column in range(NAME_COL, last + 1):
        cell = ws.cell(HEADER_ROW, column)
        cell.font = _font(bold=True)
        cell.border = RULE
        cell.alignment = CENTER if column == STATUS_COL else Alignment(vertical="center")
        if column <= STATUS_COL:
            cell.fill = _fill(HEADER)
    note = (
        "Done, In progress, and Future come from roadmap.yaml. "
        "Epic marks a section. Milestone marks a single date. "
        "Green cells are the sprints a task covers. "
        "A blue cell is the sprint that contains a milestone."
    )
    status.comment = Comment(note, "Cool Spine", width=260, height=90)


def _write_rows(ws, rows: list[dict], sprints: list[date], step: int, current: date | None) -> None:
    for offset, item in enumerate(rows):
        excel_row = HEADER_ROW + 1 + offset
        ws.row_dimensions[excel_row].height = 18
        _write_name(ws, excel_row, item)
        _write_status(ws, excel_row, item)
        for index, sprint in enumerate(sprints):
            cell = ws.cell(excel_row, FIRST_SPRINT_COL + index)
            cell.font = _font()
            cell.alignment = CENTER
            paint = _sprint_fill(item, sprint, step, current)
            if paint is not None:
                cell.fill = paint


def _write_name(ws, excel_row: int, item: dict) -> None:
    cell = ws.cell(excel_row, NAME_COL, item["name"])
    cell.alignment = NAME
    if item["kind"] == "epic":
        cell.font = _font(bold=True)
        cell.fill = _fill(EPIC)
    else:
        cell.font = _font()


def _write_status(ws, excel_row: int, item: dict) -> None:
    cell = ws.cell(excel_row, STATUS_COL, item["status"])
    cell.alignment = CENTER
    if item["kind"] == "epic":
        cell.font = _font(bold=True)
        cell.fill = _fill(EPIC)
        return
    background, color = STATUS_COLORS[item["status"]]
    cell.font = _font(10, bold=True, color=color)
    cell.fill = _fill(background)


def _finish_sheet(ws, sprint_count: int, row_count: int) -> None:
    last_col = get_column_letter(FIRST_SPRINT_COL + sprint_count - 1)
    last_row = HEADER_ROW + row_count
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last_row}"
    ws.print_title_rows = "1:4"
    ws.print_title_cols = "A:B"
    ws.print_area = f"A1:{last_col}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.tabColor = BAR
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = "Generated from roadmap.yaml"
    ws.oddFooter.right.text = "Green = task bar    Blue = milestone"
    if row_count:
        check = DataValidation(
            type="list",
            formula1='"' + ",".join(STATUS_LIST) + '"',
            allow_blank=True,
        )
        check.error = "Choose a status from the list."
        check.errorTitle = "Invalid status"
        ws.add_data_validation(check)
        check.add(f"B{HEADER_ROW + 1}:B{last_row}")


def _sprint_fill(item: dict, sprint: date, step: int, current: date | None):
    if item["kind"] == "epic":
        return _fill(EPIC)
    if _covers(item, sprint, step):
        return _fill(MILESTONE if item["kind"] == "milestone" else BAR)
    if sprint == current:
        return _fill(CURRENT)
    return None


def _covers(item: dict, sprint: date, step: int) -> bool:
    if item["kind"] == "task" and item.get("start") is not None:
        return item["start"] <= sprint <= item["end"]
    if item["kind"] == "milestone" and item.get("date") is not None:
        return sprint <= item["date"] <= sprint + timedelta(days=step - 1)
    return False


def _current_sprint(sprints: list[date], step: int, today: date) -> date | None:
    for sprint in sprints:
        if sprint <= today <= sprint + timedelta(days=step - 1):
            return sprint
    return None


def _runs(labels: list[str]) -> list[tuple[str, int, int]]:
    found: list[tuple[str, int, int]] = []
    for index, label in enumerate(labels):
        if found and found[-1][0] == label:
            found[-1] = (label, found[-1][1], index)
        else:
            found.append((label, index, index))
    return found


def _sheet_title(title: str) -> str:
    cleaned = "".join(" " if char in ":\\/?*[]" else char for char in title)
    cleaned = " ".join(cleaned.split())
    return (cleaned or "Roadmap")[:31]


def _font(size: int = 11, bold: bool = False, color: str = INK) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color)


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _parse_today(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Use YYYY-MM-DD.") from exc


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Cool Spine GANTT workbook from roadmap.yaml")
    parser.add_argument("--file", type=Path, default=HERE / "roadmap.yaml")
    parser.add_argument("--output", type=Path, default=None, help="Workbook path. Default includes today's date")
    parser.add_argument("--today", type=_parse_today, default=None, help="Date used to mark the current sprint")
    args = parser.parse_args()
    output = args.output or HERE / f"Cool_Spine_Roadmap_{date.today().isoformat()}.xlsx"
    if output.suffix.lower() != ".xlsx":
        raise SystemExit("Output must be a .xlsx file.")
    try:
        data = load_roadmap(args.file)
    except RoadmapError as exc:
        raise SystemExit(str(exc)) from exc
    today = args.today or date.today()
    save_workbook(build_workbook(data, today), output)
    tasks = sum(item["kind"] == "task" for epic in data["epics"] for item in epic["items"])
    print(f"Wrote {output} ({tasks} tasks, {data['timeline']['sprint_count']} sprints)")


if __name__ == "__main__":
    main()
